"""Parse file upload (xlsx/csv) của Analysis Job thành list[dict].

Dùng openpyxl (xlsx) và csv chuẩn. PDF chưa hỗ trợ parse cấu trúc ở MVP
(sẽ bổ sung ở Phase 4 nếu cần — bank statement PDF).
"""

import csv
import io

import frappe


def parse_input_files(job) -> list:
    """Parse tất cả file 'File Upload' trong job.input_files. Cập nhật row_count/parsed.

    Returns: list[dict] gộp từ mọi file (mỗi dict thêm khóa __source_file).
    """
    rows = []
    for src in job.input_files:
        if src.source_type != "File Upload" or not src.file:
            continue
        try:
            content = _get_file_content(src.file)
            fname = (src.file_name or src.file or "").lower()
            if fname.endswith(".csv"):
                parsed = _parse_csv(content)
            elif fname.endswith((".xlsx", ".xls")):
                parsed = _parse_xlsx(content)
            else:
                src.parse_error = f"Định dạng chưa hỗ trợ: {fname}"
                src.parsed = 0
                continue
            for r in parsed:
                r["__source_file"] = src.file_name or src.file
            rows.extend(parsed)
            src.row_count = len(parsed)
            src.parsed = 1
            src.parse_error = None
        except Exception as e:
            src.parsed = 0
            src.parse_error = str(e)[:140]
            frappe.log_error(frappe.get_traceback(), "AIO file_parser")
    return rows


def _get_file_content(file_url):
    """Lấy nội dung file (bytes) từ File DocType theo file_url."""
    file_doc = frappe.get_doc("File", {"file_url": file_url})
    return file_doc.get_content()


def _parse_csv(content) -> list:
    if isinstance(content, bytes):
        content = content.decode("utf-8-sig", errors="replace")
    reader = csv.DictReader(io.StringIO(content))
    return [dict(row) for row in reader]


def _parse_xlsx(content) -> list:
    from openpyxl import load_workbook

    if isinstance(content, str):
        content = content.encode()
    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    ws = wb.active
    rows_iter = ws.iter_rows(values_only=True)
    try:
        header = [str(h).strip() if h is not None else f"col_{i}"
                  for i, h in enumerate(next(rows_iter))]
    except StopIteration:
        return []
    out = []
    for row in rows_iter:
        if all(c is None for c in row):
            continue
        out.append({header[i]: row[i] for i in range(min(len(header), len(row)))})
    wb.close()
    return out
